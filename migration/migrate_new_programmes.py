#!/usr/bin/env python3
"""Migrates the ~114 programmes present in the staging workbook but not yet on the
website into Supabase (programmes, programme_tiers, programme_features).

Unlike import_workbook.py (which assumes an empty table), this script:
  1. Reads what's already in the `programmes` table.
  2. Matches each workbook row to an existing DB row by company (falling back to a
     name-similarity check when a company has more than one programme), so
     already-migrated programmes are never re-inserted.
  3. Inserts only the workbook rows that don't match anything already on the site.
  4. Leaves the 30 existing programmes and their tiers/features untouched.

Usage: python3 migrate_new_programmes.py [--dry-run]
"""

import sys
import difflib
from collections import defaultdict

import openpyxl
import requests

SUPABASE_URL = "https://dyuflyhkanmczwshmbyh.supabase.co"
SUPABASE_KEY = "sb_publishable_M3BltV-qjx1gED6_ktciuw_9FhwK4G5"
WORKBOOK_PATH = "/Users/verasousateixeira/Desktop/Details-Loyalty-Hub/migration/Details_Loyalty_Staging_Workbook.xlsx"

HEADERS = {
    "apikey": SUPABASE_KEY,
    "Authorization": f"Bearer {SUPABASE_KEY}",
    "Content-Type": "application/json",
    "Prefer": "return=representation",
}


def rest(table):
    return f"{SUPABASE_URL}/rest/v1/{table}"


def split_multi(val):
    if not val:
        return []
    return [p.strip() for p in str(val).split(";") if p.strip()]


def yes_no_to_bool(val):
    if val == "Yes":
        return True
    if val == "No":
        return False
    return None


def strip_sub_industry_prefix(val):
    if not val:
        return None
    if ":" in val:
        return val.split(":", 1)[1].strip()
    return val.strip()


def norm(s):
    return (s or "").strip().lower()


def load_workbook_rows():
    wb = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)

    ws = wb["PROGRAMMES"]
    headers = [c.value for c in ws[4]]
    col = headers.index
    prog_rows = [r for r in ws.iter_rows(min_row=5, values_only=True) if r[0]]

    ws2 = wb["TIERS"]
    theaders = [c.value for c in ws2[4]]
    tcol = theaders.index
    tier_rows = [r for r in ws2.iter_rows(min_row=5, values_only=True) if r[0]]

    ws3 = wb["FEATURES"]
    feat_rows = [r for r in ws3.iter_rows(min_row=5, values_only=True) if r[0]]

    programmes = []
    for r in prog_rows:
        programmes.append({
            "_row_name": r[col("Programme Name *")],
            "_row_company": r[col("Company *")],
            "programme_name": r[col("Programme Name *")],
            "company": r[col("Company *")],
            "parent_company": r[col("Parent / Main Company")],
            "cover_image_url": r[col("Logo / Image URL")],
            "launch_year": int(r[col("Launch Year")]) if r[col("Launch Year")] else None,
            "industry": r[col("Industry *")],
            "sub_industry": strip_sub_industry_prefix(r[col("Sub-Industry *")]),
            "programme_positioning": r[col("Programme Positioning")],
            "target_customer": split_multi(r[col("Target Customer")]),
            "country": r[col("Company Country *")],
            "geographic_scope": split_multi(r[col("Geographic Scope")]),
            "membership_type": r[col("Membership Type *")],
            "access_registration": r[col("Access / Registration *")],
            "mechanisms": split_multi(r[col("Mechanisms")]),
            "benefits": split_multi(r[col("Benefits")]),
            "points_expires": yes_no_to_bool(r[col("Points: Expires?")]),
            "points_expiration_period": r[col("Points: Expiration Period")],
            "points_notes": r[col("Points: Earning / Redemption Notes")],
            "discount_types": split_multi(r[col("Discount Type")]),
            "partner_companies": split_multi(r[col("Partner Companies")]),
            "source_url": r[col("Source / Programme URL")],
            "created_by": "Migration",
        })

    tiers_by_programme = {}
    for r in tier_rows:
        name = r[tcol("Programme Name *")]
        tiers_by_programme.setdefault(name, []).append({
            "tier_order": int(r[tcol("Tier Order *")]) if r[tcol("Tier Order *")] else None,
            "tier_name": r[tcol("Tier Name *")],
            "tier_price": float(r[tcol("Fee")]) if r[tcol("Fee")] not in (None, "") else None,
            "currency": r[tcol("Currency")] or "EUR",
            "qualification_amount": float(r[tcol("Qualification Amount")]) if r[tcol("Qualification Amount")] not in (None, "") else None,
            "qualification_unit": r[tcol("Qualification Unit")],
            "note": r[tcol("Note")],
        })

    features_by_programme = {}
    for r in feat_rows:
        name = r[0]
        features_by_programme.setdefault(name, []).append({"feature_name": r[1]})

    return programmes, tiers_by_programme, features_by_programme


def fetch_existing_programmes():
    resp = requests.get(rest("programmes"), headers=HEADERS,
                         params={"select": "programme_name,company"})
    resp.raise_for_status()
    return resp.json()


def match_existing(programmes, existing):
    """Returns the set of workbook row indices that already exist on the site."""
    db_by_company = defaultdict(list)
    for d in existing:
        db_by_company[norm(d["company"])].append(d)

    used_db = set()
    matched_idx = set()

    for i, p in enumerate(programmes):
        candidates = db_by_company.get(norm(p["_row_company"]), [])
        best, best_score = None, 0
        for d in candidates:
            key = (d["company"], d["programme_name"])
            if key in used_db:
                continue
            score = difflib.SequenceMatcher(None, norm(p["_row_name"]), norm(d["programme_name"])).ratio()
            if score > best_score:
                best_score, best = score, d
        if best and (best_score > 0.35 or len(candidates) == 1):
            matched_idx.add(i)
            used_db.add((best["company"], best["programme_name"]))

    return matched_idx


def main():
    dry_run = "--dry-run" in sys.argv

    programmes, tiers_by_programme, features_by_programme = load_workbook_rows()
    existing = fetch_existing_programmes()
    matched_idx = match_existing(programmes, existing)

    new_programmes = [p for i, p in enumerate(programmes) if i not in matched_idx]

    print(f"Workbook: {len(programmes)} programmes total.")
    print(f"Already on site (matched, will be skipped): {len(matched_idx)}")
    print(f"New to insert: {len(new_programmes)}")
    print()

    if dry_run:
        for p in new_programmes:
            print(f"  WOULD INSERT: {p['_row_name']} ({p['_row_company']})")
        print(f"\nDry run — no changes made. {len(new_programmes)} programmes would be inserted, "
              f"final total would be {len(existing) + len(new_programmes)}.")
        return

    inserted, failed = [], []

    for p in new_programmes:
        name = p.pop("_row_name")
        p.pop("_row_company")
        resp = requests.post(rest("programmes"), headers=HEADERS, json=p)
        if resp.status_code not in (200, 201):
            print(f"FAILED programme '{name}': {resp.status_code} {resp.text}")
            failed.append(name)
            continue
        prog_id = resp.json()[0]["id"]
        inserted.append(name)
        print(f"OK  programme '{name}' -> {prog_id}")

        tiers = tiers_by_programme.get(name, [])
        if tiers:
            for t in tiers:
                t["programme_id"] = prog_id
            tresp = requests.post(rest("programme_tiers"), headers=HEADERS, json=tiers)
            if tresp.status_code not in (200, 201):
                print(f"    FAILED tiers for '{name}': {tresp.status_code} {tresp.text}")
            else:
                print(f"    + {len(tiers)} tier(s)")

        feats = features_by_programme.get(name, [])
        if feats:
            for f in feats:
                f["programme_id"] = prog_id
            fresp = requests.post(rest("programme_features"), headers=HEADERS, json=feats)
            if fresp.status_code not in (200, 201):
                print(f"    FAILED features for '{name}': {fresp.status_code} {fresp.text}")
            else:
                print(f"    + {len(feats)} feature(s)")

    print(f"\nDone. Inserted {len(inserted)}/{len(new_programmes)} new programmes.")
    print(f"Site total should now be {len(existing) + len(inserted)}.")
    if failed:
        print("Failed:", failed)
        sys.exit(1)


if __name__ == "__main__":
    main()
