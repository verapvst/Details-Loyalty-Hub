#!/usr/bin/env python3
"""One-time addition: 21 new loyalty/membership programmes requested 2026-09-26,
closing three specific gaps identified in the benchmark sample:
  - Golf (9): private/reciprocal club networks (Eligo, Troon Privé, PlatinumClubNet)
    alongside Algarve/international private golf clubs — the sample had no
    private/reciprocal-club mechanics in Golf.
  - Food & Beverage (7) + Middle East (2): premium/luxury private members' clubs
    (JNcQUOI, Soho House, Annabel's, Zero Bond, The Supper Club / Helm, Neera,
    Capital Club Dubai) plus two subscription/curation programmes (Wine Access
    Michelin Club, Gourmet Escapes) — F&B had 0% Premium/Luxury coverage.
  - Leisure & Entertainment (1): Yellowstone Club — this industry also had 0%
    Luxury coverage.
  - Geographic diversity (2): Rotana DISCOVERY and Zomato Gold — chosen to add
    Middle East/Africa and India/Asia programmes outside the
    US/UK/Portugal/Western-Europe cluster that dominates the existing sample.

Uses the current v3 taxonomy (11 mechanisms, 9 industries + Other — see
assets/options.js / migration/remap_mechanisms_v3.md). Field values were
assigned from general knowledge plus a round of web research done for this
batch (official membership/programme pages where they exist; Wikipedia or
press coverage for a few very private clubs with no public membership page,
e.g. Queenwood). Source URLs and logos were given directly by the user for
some fields (logo, industry/country for Rotana + Zomato); everything else
(source URL where the user's own text lost its hyperlink, mechanisms,
positioning, target customer, geographic scope) is this script author's
confident first pass, not verified fact — spot-check before citing any of
these in the thesis. launch_year is left blank throughout rather than guessed.

This script does two things, in order:
  1. Appends the 21 new programmes as new PROGRAMMES rows in
     migration/Details_Loyalty_Staging_Workbook.xlsx.
  2. Inserts the same 21 programmes into Supabase (`programmes` table).

Usage: python3 add_new_programmes_2026_batch2.py [--dry-run] [--excel-only] [--db-only]
"""

import os
import sys
import openpyxl
import requests

SUPABASE_URL = "https://dyuflyhkanmczwshmbyh.supabase.co"
SUPABASE_KEY = "sb_publishable_M3BltV-qjx1gED6_ktciuw_9FhwK4G5"
WORKBOOK_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Details_Loyalty_Staging_Workbook.xlsx")

HEADERS = {
    "apikey": SUPABASE_KEY,
    "Authorization": f"Bearer {SUPABASE_KEY}",
    "Content-Type": "application/json",
    "Prefer": "return=representation",
}

# ---------------- The 21 new programmes (v3 taxonomy) ----------------
# Columns match the PROGRAMMES sheet / programmes table exactly.
# mechanisms / target_customer / geographic_scope are lists -> joined with "; " for Excel.

NEW_PROGRAMMES = [
    # ---- Golf (9) — closes the "no private/reciprocal clubs" gap ----
    dict(programme_name="Vilamoura Golf Membership", company="Vilamoura Golf", parent_company=None,
         cover_image_url="https://www.vilamouragolf.com/wp-content/uploads/2025/02/VM-Golf-Logo-Slate-Green-RGB.png",
         industry="Golf", programme_positioning="Premium",
         target_customer=["Enthusiasts / Hobbyists", "High-Value Customers"], country="Portugal",
         geographic_scope=["Portugal"], membership_type="Paid", access_registration="Application required",
         mechanisms=["Member pricing", "Privileged access", "Included member benefits"],
         source_url="https://www.vilamouragolf.com/"),
    dict(programme_name="Vale do Lobo Golf Club Membership", company="Vale do Lobo Golf Club", parent_company="Vale do Lobo Resort",
         cover_image_url="https://encrypted-tbn0.gstatic.com/images?q=tbn:ANd9GcScsGWEwTQ0yq4jNdzWjJ3NTOHbAygD1i-Aeoe3MIPiRKmfqTULFowr-fsI&s=10",
         industry="Golf", programme_positioning="Luxury",
         target_customer=["Premium / Luxury Customers", "Enthusiasts / Hobbyists"], country="Portugal",
         geographic_scope=["Portugal"], membership_type="Paid", access_registration="Application required",
         mechanisms=["Member pricing", "Privileged access", "Included member benefits"],
         source_url="https://www.valedolobo.com/wp-content/uploads/2025/01/2025_Golf_Club_Membership.pdf"),
    dict(programme_name="Amendoeira Golf Resort Membership", company="Amendoeira Golf Resort", parent_company=None,
         cover_image_url="https://encrypted-tbn0.gstatic.com/images?q=tbn:ANd9GcT4sVOPTInd3Y_LcwZgd0TAZl-NQt6wnPONehLhgbBicyYowKJW9kgBRmY&s=10",
         industry="Golf", programme_positioning="Premium",
         target_customer=["Enthusiasts / Hobbyists"], country="Portugal",
         geographic_scope=["Portugal"], membership_type="Paid", access_registration="Open registration",
         mechanisms=["Member pricing", "Privileged access"],
         source_url="https://www.amendoeiraresort.com/en/golf/"),
    dict(programme_name="Monte Rei Golf & Country Club", company="Monte Rei Golf & Country Club", parent_company=None,
         cover_image_url="https://encrypted-tbn0.gstatic.com/images?q=tbn:ANd9GcQIx2SWFzivpIg-WNpI4Ptfpj0mXl2cn4IwyAtv778VwvRzmA51SmdFsxug&s=10",
         industry="Golf", programme_positioning="Luxury",
         target_customer=["Premium / Luxury Customers", "Enthusiasts / Hobbyists"], country="Portugal",
         geographic_scope=["Portugal"], membership_type="Paid", access_registration="Application required",
         mechanisms=["Privileged access", "Member experiences & events", "Community", "Included member benefits"],
         source_url="https://www.monte-rei.com/en/golf/"),
    dict(programme_name="Eligo Club", company="Eligo Club", parent_company=None,
         cover_image_url="https://p-gpb8fhd4b9fbh6fy.z01.azurefd.net/logos/209086c0-e6fe-4498-877b-964cac90e4ab/medium.png",
         industry="Golf", programme_positioning="Premium",
         target_customer=["Enthusiasts / Hobbyists", "High-Value Customers"], country="United States",
         geographic_scope=["Europe", "North America"], membership_type="Paid", access_registration="Application required",
         mechanisms=["External partner network", "Privileged access", "Member experiences & events", "Community"],
         source_url="https://www.eligoclub.com/"),
    dict(programme_name="Troon Privé (Private Club Golf)", company="Troon", parent_company=None,
         cover_image_url="https://www.oakspga.cz/wp-content/uploads/sites/9225/2023/08/TroonPrivilegesLogo-1.png?w=1024",
         industry="Golf", programme_positioning="Premium",
         target_customer=["Enthusiasts / Hobbyists", "High-Value Customers"], country="United States",
         geographic_scope=["Global"], membership_type="Paid", access_registration="Application required",
         mechanisms=["External partner network", "Privileged access", "Member pricing"],
         source_url="https://troon.com/member-programs/private-club-golf"),
    dict(programme_name="PlatinumClubNet", company="PlatinumClubNet", parent_company="MobiCom America Inc.",
         cover_image_url="https://www.platinumclubnet.com/_next/image?url=%2F_next%2Fstatic%2Fmedia%2Flog1.73b78c22.png&w=256&q=75",
         industry="Golf", programme_positioning="Luxury",
         target_customer=["High-Value Customers", "Premium / Luxury Customers"], country="United States",
         geographic_scope=["Global"], membership_type="Other", access_registration="Invitation only",
         mechanisms=["External partner network", "Privileged access", "Community"],
         source_url="https://www.platinumclubnet.com/"),
    dict(programme_name="Les Bordes Golf Club", company="Les Bordes Golf Club", parent_company=None,
         cover_image_url="https://golfcoursegurus.com/reviews/images/logos/Les-Bordes-Golf-Club-New.png",
         industry="Golf", programme_positioning="Luxury",
         target_customer=["Premium / Luxury Customers", "Enthusiasts / Hobbyists"], country="France",
         geographic_scope=["Europe"], membership_type="Paid", access_registration="Application required",
         mechanisms=["Privileged access", "Community", "Member experiences & events", "Included member benefits"],
         source_url="https://www.lesbordesgolfclub.com/club/membership/"),
    dict(programme_name="Queenwood Golf Club", company="Queenwood Golf Club", parent_company=None,
         cover_image_url="https://media.licdn.com/dms/image/v2/C560BAQH6rhMQDhVYgw/company-logo_200_200/company-logo_200_200/0/1630610259895/queenwood_golf_club_limited_logo?e=2147483647&v=beta&t=5UeZPLDRJIY2irAJfir_jEahHJvCVRwAHD_n3hifeSk",
         industry="Golf", programme_positioning="Luxury",
         target_customer=["Premium / Luxury Customers", "High-Value Customers"], country="United Kingdom",
         geographic_scope=["Europe"], membership_type="Paid", access_registration="Invitation only",
         mechanisms=["Privileged access", "Community", "Member experiences & events"],
         source_url="https://en.wikipedia.org/wiki/Queenwood_Golf_Club"),

    # ---- Food & Beverage (7) — closes the 0%-Premium/Luxury gap ----
    dict(programme_name="JNcQUOI Club", company="JNcQUOI", parent_company="Amorim Luxury",
         cover_image_url="https://www.visitgrandola.com/cmgrandola/uploads/poi/image/995/jncquoi.JPG",
         industry="Food & Beverage", programme_positioning="Luxury",
         target_customer=["Premium / Luxury Customers", "International Customers"], country="Portugal",
         geographic_scope=["Portugal"], membership_type="Paid", access_registration="Application required",
         mechanisms=["Member pricing", "Privileged access", "Included member benefits", "Member experiences & events"],
         source_url="https://www.jncquoi.com/en/membership/"),
    dict(programme_name="Soho House", company="Soho House & Co", parent_company=None,
         cover_image_url="https://i.pinimg.com/736x/d3/94/ca/d394ca792d0594d99eb16795dfa0c73a.jpg",
         industry="Food & Beverage", programme_positioning="Luxury",
         target_customer=["Premium / Luxury Customers", "International Customers"], country="United Kingdom",
         geographic_scope=["Global"], membership_type="Subscription", access_registration="Application required",
         mechanisms=["Community", "Privileged access", "Member experiences & events", "Included member benefits", "Ecosystem cross-use"],
         source_url="https://www.sohohouse.com/"),
    dict(programme_name="Annabel's (Birley Clubs)", company="Annabel's", parent_company="Birley Clubs",
         cover_image_url="https://storage.googleapis.com/talentfunnel-cms-career-site-storage/birley-clubs/Annabels_2_842f94c806_bfb48f2e97/Annabels_2_842f94c806_bfb48f2e97.svg",
         industry="Food & Beverage", programme_positioning="Luxury",
         target_customer=["Premium / Luxury Customers"], country="United Kingdom",
         geographic_scope=["Europe"], membership_type="Paid", access_registration="Invitation only",
         mechanisms=["Privileged access", "Member experiences & events", "Community", "Included member benefits"],
         source_url="https://annabels.co.uk/"),
    dict(programme_name="Zero Bond", company="Zero Bond", parent_company="Bond Hospitality",
         cover_image_url="https://encrypted-tbn0.gstatic.com/images?q=tbn:ANd9GcT1GOsXSa2VYZkjB_ezGUDz61tRPScQqS6q0nMx4RgKu4vfaCnbFDoh7hs&s=10",
         industry="Food & Beverage", programme_positioning="Luxury",
         target_customer=["Premium / Luxury Customers", "High-Value Customers"], country="United States",
         geographic_scope=["North America"], membership_type="Paid", access_registration="Application required",
         mechanisms=["Privileged access", "Community", "Member experiences & events"],
         source_url="https://zerobondny.com/apply-for-membership/"),
    dict(programme_name="The Supper Club", company="The Supper Club (now Helm)", parent_company=None,
         cover_image_url="https://mma.prnewswire.com/media/2794987/supper_club_Logo.jpg?w=300",
         industry="Food & Beverage", programme_positioning="Premium",
         target_customer=["Business Customers", "High-Value Customers"], country="United Kingdom",
         geographic_scope=["Europe"], membership_type="Subscription", access_registration="Application required",
         mechanisms=["Community", "Member experiences & events", "Privileged access"],
         source_url="https://www.helmclub.co/"),
    dict(programme_name="Wine Access Michelin Club", company="Wine Access", parent_company=None,
         cover_image_url="https://d3h1lg3ksw6i6b.cloudfront.net/media/image/2022/02/03/8f721ef882bc48ad93a33c92171094fa_wine+access+vertical+logo+030222.png",
         industry="Food & Beverage", programme_positioning="Premium",
         target_customer=["Enthusiasts / Hobbyists", "High-Value Customers"], country="United States",
         geographic_scope=["North America", "Global"], membership_type="Subscription", access_registration="Open registration",
         mechanisms=["Member pricing", "Included member benefits", "External partner network", "Privileged access"],
         source_url="https://www.wineaccess.com/subscriptions/"),
    dict(programme_name="Gourmet Escapes Privilege Club", company="Gourmet Escapes", parent_company=None,
         cover_image_url="https://gourmet-escapes.com/wp-content/uploads/2017/09/logo.png",
         industry="Food & Beverage", programme_positioning="Premium",
         target_customer=["Premium / Luxury Customers", "Enthusiasts / Hobbyists"], country="United Kingdom",
         geographic_scope=["Europe"], membership_type="Subscription", access_registration="Open registration",
         mechanisms=["Member pricing", "External partner network", "Included member benefits", "Privileged access"],
         source_url="https://www.gourmet-escapes.com/privilege-club/"),

    # ---- Leisure & Entertainment (1) — this industry also had zero Luxury ----
    dict(programme_name="Yellowstone Club", company="Yellowstone Club", parent_company="CrossHarbor Capital Partners",
         cover_image_url="https://yellowstoneclub.com/logo-dark.svg",
         industry="Leisure & Entertainment", programme_positioning="Luxury",
         target_customer=["Premium / Luxury Customers", "High-Value Customers"], country="United States",
         geographic_scope=["North America"], membership_type="Paid", access_registration="Invitation only",
         mechanisms=["Privileged access", "Community", "Member experiences & events", "Included member benefits"],
         source_url="https://yellowstoneclub.com/"),

    # ---- Geographic diversity: Middle East (2) ----
    dict(programme_name="Neera (Al Habtoor City, Dubai)", company="Neera Private Members Club", parent_company="Al Habtoor Group",
         cover_image_url="https://neeradubai.com/images/neera/peacock-feather.png",
         industry="Food & Beverage", programme_positioning="Luxury",
         target_customer=["Premium / Luxury Customers", "High-Value Customers"], country="United Arab Emirates",
         geographic_scope=["Middle East"], membership_type="Paid", access_registration="Application required",
         mechanisms=["Included member benefits", "Member experiences & events", "Privileged access", "Community"],
         source_url="https://neeradubai.com/"),
    dict(programme_name="Capital Club Dubai", company="Capital Club Dubai", parent_company=None,
         cover_image_url="https://capitalclubdubai.com/wp-content/uploads/2021/04/cropped-Fatma-Al-Mahmoud-04.png",
         industry="Food & Beverage", programme_positioning="Luxury",
         target_customer=["Business Customers", "High-Value Customers"], country="United Arab Emirates",
         geographic_scope=["Middle East"], membership_type="Paid", access_registration="Invitation only",
         mechanisms=["Privileged access", "External partner network", "Community", "Member experiences & events"],
         source_url="https://capitalclubdubai.com/membership/"),

    # ---- Geographic diversity: outside US/UK/Portugal/Western-Europe cluster ----
    dict(programme_name="Rotana DISCOVERY", company="Rotana Hotels & Resorts", parent_company=None,
         cover_image_url="https://cdb.rotana.com/rcnimagelib/nx_logo_w_10.png",
         industry="Hotels & Hospitality", programme_positioning="Premium",
         target_customer=["Business Customers", "International Customers"], country="United Arab Emirates",
         geographic_scope=["Middle East", "Africa"], membership_type="Free", access_registration="Open registration",
         mechanisms=["Spend-based earning", "Earned status", "External partner network", "Privileged access"],
         source_url="https://www.rotana.com/rotanadiscovery"),
    dict(programme_name="Zomato Gold", company="Zomato", parent_company="Eternal Limited",
         cover_image_url="https://b.zmtcdn.com/data/o2_assets/a8d9789aef91b9ffbca4f613010bb8d21714487350.png",
         industry="Food & Beverage", programme_positioning="Premium",
         target_customer=["Mass Market", "Frequent Customers"], country="India",
         geographic_scope=["Asia"], membership_type="Subscription", access_registration="Open registration",
         mechanisms=["Member pricing", "Included member benefits", "Privileged access"],
         source_url="https://www.zomato.com/gold"),
]


def join_multi(vals):
    return "; ".join(vals) if vals else None


def update_excel(dry_run=False):
    wb = openpyxl.load_workbook(WORKBOOK_PATH, data_only=False)
    ws = wb["PROGRAMMES"]

    next_row = ws.max_row + 1
    for p in NEW_PROGRAMMES:
        ws.cell(row=next_row, column=1).value = p["programme_name"]
        ws.cell(row=next_row, column=2).value = p["company"]
        ws.cell(row=next_row, column=3).value = p.get("parent_company")
        ws.cell(row=next_row, column=4).value = p.get("cover_image_url")
        ws.cell(row=next_row, column=5).value = p.get("launch_year")
        ws.cell(row=next_row, column=6).value = p["industry"]
        ws.cell(row=next_row, column=7).value = None  # Sub-Industry — deprecated
        ws.cell(row=next_row, column=8).value = p.get("programme_positioning")
        ws.cell(row=next_row, column=9).value = join_multi(p.get("target_customer"))
        ws.cell(row=next_row, column=10).value = p["country"]
        ws.cell(row=next_row, column=11).value = join_multi(p.get("geographic_scope"))
        ws.cell(row=next_row, column=12).value = p["membership_type"]
        ws.cell(row=next_row, column=13).value = p["access_registration"]
        ws.cell(row=next_row, column=14).value = join_multi(p.get("mechanisms"))
        ws.cell(row=next_row, column=15).value = None  # Benefits — deprecated
        ws.cell(row=next_row, column=21).value = p["source_url"]
        next_row += 1
    print(f"Appended {len(NEW_PROGRAMMES)} new programme row(s) (rows {ws.max_row - len(NEW_PROGRAMMES) + 1}-{ws.max_row}).")

    if dry_run:
        print("--dry-run: workbook not saved.")
    else:
        wb.save(WORKBOOK_PATH)
        print(f"Saved {WORKBOOK_PATH}")


def insert_into_supabase(dry_run=False):
    if dry_run:
        print(f"--dry-run: would insert {len(NEW_PROGRAMMES)} programme(s) into Supabase.")
        return
    ok, failed = 0, []
    for p in NEW_PROGRAMMES:
        payload = {
            "programme_name": p["programme_name"],
            "company": p["company"],
            "parent_company": p.get("parent_company"),
            "cover_image_url": p.get("cover_image_url"),
            "launch_year": p.get("launch_year"),
            "industry": p["industry"],
            "programme_positioning": p.get("programme_positioning"),
            "target_customer": p.get("target_customer") or [],
            "country": p["country"],
            "geographic_scope": p.get("geographic_scope") or [],
            "membership_type": p["membership_type"],
            "access_registration": p["access_registration"],
            "mechanisms": p.get("mechanisms") or [],
            "source_url": p["source_url"],
            "created_by": "Migration",
        }
        resp = requests.post(f"{SUPABASE_URL}/rest/v1/programmes", headers=HEADERS, json=payload)
        if resp.status_code not in (200, 201):
            print(f"FAILED '{p['programme_name']}': {resp.status_code} {resp.text}")
            failed.append(p["programme_name"])
            continue
        ok += 1
        print(f"OK  '{p['programme_name']}' -> {resp.json()[0]['id']}")
    print(f"\nInserted {ok}/{len(NEW_PROGRAMMES)} programme(s) into Supabase.")
    if failed:
        print("Failed:", failed)


if __name__ == "__main__":
    dry_run = "--dry-run" in sys.argv
    do_excel = "--db-only" not in sys.argv
    do_db = "--excel-only" not in sys.argv
    if do_excel:
        update_excel(dry_run=dry_run)
    if do_db:
        insert_into_supabase(dry_run=dry_run)
