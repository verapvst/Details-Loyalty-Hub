#!/usr/bin/env python3
"""One-time import: reads the filled staging workbook and inserts everything into
Supabase (programmes, programme_tiers, programme_features), matching the exact
column names the live app uses.

Usage: python3 import_workbook.py
"""

import sys
import openpyxl
import requests

SUPABASE_URL = "https://dyuflyhkanmczwshmbyh.supabase.co"
SUPABASE_KEY = "sb_publishable_M3BltV-qjx1gED6_ktciuw_9FhwK4G5"
WORKBOOK_PATH = "/Users/verasousateixeira/Desktop/Details-Loyalty-Hub/migration/Details_Loyalty_Staging_Workbook.xlsx"

HEADERS = {
    "apikey": SUPABASE_KEY,
    "Authorization": f"Bearer {SUPABASE_KEY}",
    "Content-Type": "application/json",
    "Prefer": "return=representation",
}


def rest(table):
    return f"{SUPABASE_URL}/rest/v1/{table}"


def split_multi(val):
    if not val:
        return []
    return [p.strip() for p in str(val).split(";") if p.strip()]


def yes_no_to_bool(val):
    if val == "Yes":
        return True
    if val == "No":
        return False
    return None


def strip_sub_industry_prefix(val):
    if not val:
        return None
    if ":" in val:
        return val.split(":", 1)[1].strip()
    return val.strip()


def load_workbook_rows():
    wb = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)

    ws = wb["PROGRAMMES"]
    headers = [c.value for c in ws[4]]
    col = headers.index
    prog_rows = [r for r in ws.iter_rows(min_row=5, values_only=True) if r[0]]

    ws2 = wb["TIERS"]
    theaders = [c.value for c in ws2[4]]
    tcol = theaders.index
    tier_rows = [r for r in ws2.iter_rows(min_row=5, values_only=True) if r[0]]

    ws3 = wb["FEATURES"]
    feat_rows = [r for r in ws3.iter_rows(min_row=5, values_only=True) if r[0]]

    programmes = []
    for r in prog_rows:
        programmes.append({
            "_row_name": r[col("Programme Name *")],
            "programme_name": r[col("Programme Name *")],
            "company": r[col("Company *")],
            "parent_company": r[col("Parent / Main Company")],
            "cover_image_url": r[col("Logo / Image URL")],
            "launch_year": int(r[col("Launch Year")]) if r[col("Launch Year")] else None,
            "industry": r[col("Industry *")],
            "sub_industry": strip_sub_industry_prefix(r[col("Sub-Industry *")]),
            "programme_positioning": r[col("Programme Positioning")],
            "target_customer": split_multi(r[col("Target Customer")]),
            "country": r[col("Company Country *")],
            "geographic_scope": split_multi(r[col("Geographic Scope")]),
            "membership_type": r[col("Membership Type *")],
            "access_registration": r[col("Access / Registration *")],
            "mechanisms": split_multi(r[col("Mechanisms")]),
            "benefits": split_multi(r[col("Benefits")]),
            "points_expires": yes_no_to_bool(r[col("Points: Expires?")]),
            "points_expiration_period": r[col("Points: Expiration Period")],
            "points_notes": r[col("Points: Earning / Redemption Notes")],
            "discount_types": split_multi(r[col("Discount Type")]),
            "partner_companies": split_multi(r[col("Partner Companies")]),
            "source_url": r[col("Source / Programme URL")],
            "created_by": "Migration",
        })

    tiers_by_programme = {}
    for r in tier_rows:
        name = r[tcol("Programme Name *")]
        tiers_by_programme.setdefault(name, []).append({
            "tier_order": int(r[tcol("Tier Order *")]) if r[tcol("Tier Order *")] else None,
            "tier_name": r[tcol("Tier Name *")],
            "tier_price": float(r[tcol("Fee")]) if r[tcol("Fee")] not in (None, "") else None,
            "currency": r[tcol("Currency")] or "EUR",
            "qualification_amount": float(r[tcol("Qualification Amount")]) if r[tcol("Qualification Amount")] not in (None, "") else None,
            "qualification_unit": r[tcol("Qualification Unit")],
            "note": r[tcol("Note")],
        })

    features_by_programme = {}
    for r in feat_rows:
        name = r[0]
        features_by_programme.setdefault(name, []).append({"feature_name": r[1]})

    return programmes, tiers_by_programme, features_by_programme


def main():
    programmes, tiers_by_programme, features_by_programme = load_workbook_rows()
    print(f"Loaded {len(programmes)} programmes, "
          f"{sum(len(v) for v in tiers_by_programme.values())} tiers, "
          f"{sum(len(v) for v in features_by_programme.values())} features from workbook.\n")

    inserted, failed = [], []

    for p in programmes:
        name = p.pop("_row_name")
        resp = requests.post(rest("programmes"), headers=HEADERS, json=p)
        if resp.status_code not in (200, 201):
            print(f"FAILED programme '{name}': {resp.status_code} {resp.text}")
            failed.append(name)
            continue
        prog_id = resp.json()[0]["id"]
        inserted.append(name)
        print(f"OK  programme '{name}' -> {prog_id}")

        tiers = tiers_by_programme.get(name, [])
        if tiers:
            for t in tiers:
                t["programme_id"] = prog_id
            tresp = requests.post(rest("programme_tiers"), headers=HEADERS, json=tiers)
            if tresp.status_code not in (200, 201):
                print(f"    FAILED tiers for '{name}': {tresp.status_code} {tresp.text}")
            else:
                print(f"    + {len(tiers)} tier(s)")

        feats = features_by_programme.get(name, [])
        if feats:
            for f in feats:
                f["programme_id"] = prog_id
            fresp = requests.post(rest("programme_features"), headers=HEADERS, json=feats)
            if fresp.status_code not in (200, 201):
                print(f"    FAILED features for '{name}': {fresp.status_code} {fresp.text}")
            else:
                print(f"    + {len(feats)} feature(s)")

    print(f"\nDone. Inserted {len(inserted)}/{len(programmes)} programmes.")
    if failed:
        print("Failed:", failed)
        sys.exit(1)


if __name__ == "__main__":
    main()
